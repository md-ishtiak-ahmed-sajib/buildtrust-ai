import csv
from io import StringIO, BytesIO
import json
from pathlib import Path
import re
import subprocess
import sys
import zipfile
from xml.etree.ElementTree import ParseError
from openpyxl import load_workbook, Workbook
from openpyxl.utils.exceptions import InvalidFileException
from .advanced_schemas import (
    QuotationInput,
    EstimateFields,
    EstimateLine,
    RateComponent,
)
from .bounded_json import loads as bounded_loads


def table_rows(contents, extension, required=()):
    try:
        return _table_rows(contents, extension, required)
    except (
        zipfile.BadZipFile,
        InvalidFileException,
        ParseError,
        StopIteration,
        KeyError,
        TypeError,
    ):
        raise ValueError(
            "Invalid or incomplete financial table; use a valid CSV/XLSX template"
        )


class TableInputError(ValueError):
    def __init__(self, message, fields):
        super().__init__(message)
        self.fields = fields


def validate_headers(headers, required):
    if not headers or len(headers) > 60 or any(not h for h in headers):
        raise ValueError("Headers require 1–60 nonempty column names")
    if len(set(headers)) != len(headers):
        raise TableInputError(
            "Duplicate column headers are not permitted",
            [{"location": ["headers"], "message": "Column names must be unique"}],
        )
    missing = sorted(set(required) - set(headers))
    if missing:
        raise TableInputError(
            "Missing required columns: " + ", ".join(missing),
            [
                {
                    "location": ["headers", column],
                    "message": "Required column is missing",
                }
                for column in missing
            ],
        )


def _table_rows(contents, extension, required):
    if extension == ".csv":
        reader = csv.DictReader(StringIO(contents.decode("utf-8-sig")))
        headers = [h.strip() for h in (reader.fieldnames or [])]
        validate_headers(headers, required)
        reader.fieldnames = headers
        rows = []
        for index, row in enumerate(reader, start=1):
            if index > 500:
                raise ValueError("Import is limited to 500 rows")
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"Row {index + 1}: cell count differs from headers")
            rows.append(row)
    elif extension == ".xlsx":
        with zipfile.ZipFile(BytesIO(contents)) as archive:
            if (
                sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024
                or len(archive.infolist()) > 200
            ):
                raise ValueError("Expanded XLSX exceeds import bounds")
        workbook = load_workbook(
            BytesIO(contents), read_only=True, data_only=False, keep_links=False
        )
        try:
            sheet = workbook.active
            if sheet.max_row > 501 or sheet.max_column > 60:
                raise ValueError("Import is limited to 500 rows and 60 columns")
            values = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(values)]
            validate_headers(headers, required)
            rows = []
            for index, row in enumerate(values, start=1):
                if index > 500:
                    raise ValueError("Import is limited to 500 rows")
                if any(isinstance(v, str) and v.startswith("=") for v in row):
                    raise ValueError(
                        "Formula cells are not accepted in financial imports"
                    )
                rows.append(
                    {k: str(v) for k, v in zip(headers, row) if k and v is not None}
                )
        finally:
            workbook.close()
    else:
        raise ValueError("Use a documented CSV or XLSX table")
    if not rows or len(rows) > 500:
        raise ValueError("Import requires 1–500 rows")
    return rows


def quotation_rows(contents, extension, material_id, metadata):
    parsed = []
    for row in table_rows(contents, extension):
        row = {k: v for k, v in row.items() if k and v not in (None, "")}
        spec = bounded_loads(row.pop("specification", "{}"))
        for key in list(row):
            if key.startswith("spec_"):
                spec[key[5:]] = row.pop(key)
        values = {
            **metadata,
            **row,
            "material_id": material_id,
            "specification": spec or metadata.get("specification", {}),
        }
        parsed.append(QuotationInput.model_validate(values))
    return parsed


def pdf_suggestions(path, settings=None):
    from .ocr import worker_args

    try:
        proc = subprocess.run(
            [
                sys.executable,
                str(Path(__file__).with_name("pdf_worker.py")),
                str(path),
                *worker_args(settings),
            ],
            capture_output=True,
            timeout=45 if settings and settings.ocr_enabled else 15,
            check=True,
        )
    except subprocess.SubprocessError:
        raise ValueError(
            "PDF extraction failed or exceeded its limits; use manual quotation entry"
        )
    result = json.loads(proc.stdout)
    text = "\n".join(result["pages"])[:100000]
    suggestions, evidence = {}, {}
    labels = {
        "supplier": "Supplier",
        "number": "Quotation Number",
        "title": "Product",
        "price": "Unit Price",
        "currency": "Currency",
        "unit": "Unit",
        "quoted_quantity": "Quantity",
        "quotation_date": "Quotation Date",
        "valid_until": "Valid Until",
    }
    for field, label in labels.items():
        found = re.search(r"^" + re.escape(label) + r"\s*:\s*(.+)$", text, re.M | re.I)
        if found:
            suggestions[field] = found.group(1).strip()
            pages = [
                i + 1
                for i, page in enumerate(result["pages"])
                if found.group(0) in page
            ]
            evidence[field] = {
                "page": pages[0] if len(pages) == 1 else None,
                "method": "PDF_LABEL_SUGGESTION",
                "confidence": None,
                "review_required": True,
            }
    return suggestions, {
        "fields": evidence,
        "text_preview": text[:12000],
        "pages": len(result["pages"]),
        "page_evidence": result.get("page_evidence", []),
        "ocr": result.get("ocr", {}),
        "notice": "Untrusted text. Suggested values need human review; scanned PDFs require manual entry.",
    }


def safe_cell(value):
    text = str(value if value is not None else "")
    return "'" + text if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text


def export_table(headers, rows, extension):
    rows = [[safe_cell(value) for value in row] for row in rows]
    if extension == "csv":
        buffer = StringIO(newline="")
        writer = csv.writer(buffer)
        writer.writerow(headers)
        writer.writerows(rows)
        return buffer.getvalue().encode("utf-8-sig"), "text/csv"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BOQ"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return (
        stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def boq_import(contents, extension):
    lines = []
    rows = table_rows(contents, extension, ("code", "description", "unit", "quantity", "unit_rate"))
    settings_keys = ("site_overhead_percent", "general_overhead_percent", "profit_percent", "tax_percent", "contingency_percent", "tax_treatment")
    context = {}
    if rows and rows[0].get("buildtrust_export_version") == "2":
        context = {key: rows[0].get(key) for key in settings_keys}
        if any({key: row.get(key) for key in settings_keys} != context or row.get("buildtrust_export_version") != "2" for row in rows):
            raise ValueError("Export allowance settings must agree across every row")
    for index, row in enumerate(rows):
        components = []
        if context:
            try:
                values = bounded_loads(row["components_json"])
                if not isinstance(values, list) or not 1 <= len(values) <= 100:
                    raise ValueError("Export requires 1–100 components per line")
                for value in values:
                    component = RateComponent.model_validate(value)
                    original = component.selection_id
                    component.material_id = component.selection_id = None
                    component.source_note = (f"Imported unverified calculation; original selection reference {original or 'manual'}. " + component.source_note)[:1000]
                    components.append(component)
            except (KeyError, TypeError, json.JSONDecodeError) as error:
                raise ValueError(f"Invalid component evidence on row {index + 1}") from error
        else:
            # Flat external rates are complete manual rates; do not apply the
            # previous draft's allowances again or classify composite work as material.
            components = [RateComponent(id=f"rate-{index + 1}", category="OTHER", description=row["description"],
                          unit=row["unit"], unit_price=row["unit_rate"], source_note=row.get("source_note") or "Imported final manual rate; reverify sources before approval")]
        lines.append(
            EstimateLine(
                id=row.get("id") or f"import-{index + 1}",
                section=row.get("section") or "Works",
                code=row["code"],
                description=row["description"],
                unit=row["unit"],
                quantity=row["quantity"],
                cost_code=row.get("cost_code", ""),
                due_date=row.get("due_date") or None,
                notes=row.get("notes", ""),
                drawing_reference=row.get("drawing_reference", ""),
                components=components,
            )
        )
    return EstimateFields(lines=lines, **context)
